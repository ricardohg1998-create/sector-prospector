import json
from pathlib import Path
import tempfile
import unittest

from sistema_prospeccion.engine import EXPORT_FIELDS
from sistema_prospeccion.excel_export import export_file


class PublicExcelTests(unittest.TestCase):
    def test_keeps_26_fields_text_identifiers_and_formula_like_values(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest('optional openpyxl is not installed')
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / 'input.json'
            record = {field: '' for field in EXPORT_FIELDS}
            record.update(name='=1+1', postal_code='01001', phone='+34000000000', email='demo@example.invalid',
                          additional_emails=['other@example.invalid'])
            source.write_text(json.dumps({'summary': {'status': '=unsafe'}, 'companies': [record],
                                          'others': [], 'pending': [], 'reserve': []}), encoding='utf-8')
            output = export_file(source, root / 'example.xlsx')
            wb = load_workbook(output)
            sheet = wb['Empresas']
            self.assertEqual(len(EXPORT_FIELDS), 26)
            self.assertEqual([cell.value for cell in sheet[1]], list(EXPORT_FIELDS))
            values = {field: sheet.cell(2, index + 1) for index, field in enumerate(EXPORT_FIELDS)}
            self.assertEqual(values['postal_code'].value, '01001')
            self.assertEqual(values['phone'].value, '+34000000000')
            self.assertEqual(values['name'].value, '=1+1')
            self.assertEqual(values['name'].data_type, 's')
            self.assertEqual(wb['Resumen']['B2'].data_type, 's')
            self.assertEqual(sheet.freeze_panes, 'A2')
            wb.close()


if __name__ == '__main__':
    unittest.main()
