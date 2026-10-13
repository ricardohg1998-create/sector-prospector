"""Portable XLSX exporter for the public extraction; no Codex runtime dependency."""
import argparse
import json
from pathlib import Path

from .engine import EXPORT_FIELDS


def export_file(input_path, output_path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError('XLSX export requires openpyxl; install requirements-excel.txt or use --csv-only') from exc
    data = json.loads(Path(input_path).read_text(encoding='utf-8'))
    wb = Workbook()
    summary = wb.active
    summary.title = 'Resumen'
    summary.append(['Field', 'Value'])
    for key, value in data.get('summary', {}).items():
        summary.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    groups = [('Empresas', 'companies'), ('Otros contactos', 'others'), ('Pendientes', 'pending'),
              ('Reserva', 'reserve'), ('Emails compartidos', 'shared')]
    for title, key in groups:
        if key not in data:
            continue
        sheet = wb.create_sheet(title)
        sheet.append(list(EXPORT_FIELDS))
        for record in data[key]:
            row = []
            for field in EXPORT_FIELDS:
                value = record.get(field, '')
                if isinstance(value, list):
                    value = ' | '.join(map(str, value))
                row.append('' if value is None else str(value))
            sheet.append(row)
            # Explicit string cells preserve leading zeroes and prevent formulas.
            for cell in sheet[sheet.max_row]:
                cell.data_type = 's'
                cell.number_format = '@'
        sheet.auto_filter.ref = sheet.dimensions
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='163D4A')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(45, max(16, len(str(column[0].value)) + 2))
        # Summary text must also remain text when it begins with formula syntax.
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    export_file(args.input, args.output)
